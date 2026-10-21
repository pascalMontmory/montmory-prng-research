"""Make a readable workbook from unchanged, independently audited CSV data."""
from collections import defaultdict
import csv
import json
from pathlib import Path
import statistics as st

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "20261021T185102Z-tpu-opt/campaign"
AUDIT = ROOT / "results/audit-20261021T185102Z-tpu-opt"
DEST = ROOT.parent / "Montmory_TPU_resultats_lisibles.xlsx"
audit = json.loads((AUDIT / "audit.json").read_text())
report = json.loads((SOURCE / "report.json").read_text())
assert audit["status"] == report["status"] == "pass"
with (SOURCE / "measurements.csv").open() as file:
    raw = list(csv.DictReader(file))
for row in raw:
    for key in ("words", "block", "seed", "repetition", "nanoseconds"):
        row[key] = int(row[key])
sizes = report["sizes"]
labels = dict(zip(sizes, ("512 Kio", "8 Mio", "64 Mio", "128 Mio")))
times = defaultdict(list)
for r in raw:
    times[r["phase"], r["words"], r["engine"], r["seed"]].append(r["nanoseconds"])


def median(count, name, seed):
    return st.median(times["confirmation", count, name, seed])


def family(name):
    if name.startswith(("jax_", "pallas_")):
        return "Montmory"
    if name.startswith("threefry"):
        return "Threefry"
    if name.startswith("philox"):
        return "Philox"
    return "RBG : même backend" if name == "rbg_same_backend" else "Écriture : NON PRNG"


def implementation(name):
    if name.startswith(("pallas_", "threefry_pallas_")):
        return "Pallas"
    return "JAX"


def note(count, name):
    selection = report["selection"][str(count)]
    if name == selection["recommended_after_confirmation"]:
        return "Montmory retenu"
    if name == "pallas_original_32":
        return "Montmory précédent"
    if family(name) == "Montmory":
        return "Montmory mesuré, non retenu"
    return "Comparateur"


wb = Workbook()
guide = wb.active
guide.title = "Lire d'abord"
guide.append(["Montmory TPU — comprendre les résultats", ""])
guide_rows = [
    ("Matériel", "TPU v5 lite, un périphérique, une session ; JAX et jaxlib 0.10.2."),
    ("Commencer ici", "Ouvrir 'Exemple 128 Mio', puis 'Synthèse'. Filtrer un seul volume avant de comparer les moteurs."),
    ("Temps (ms)", "Plus petit = plus rapide. 1 000 000 nanosecondes = 1 milliseconde."),
    ("Débit (Go/s)", "Plus grand = plus rapide. Un mot vaut 8 octets. Débit = mots × 8 / nanosecondes. Go/s décimaux = GB/s du rapport."),
    ("Volume", "512 Kio = 65 536 mots ; 8 Mio = 1 048 576 ; 64 Mio = 8 388 608 ; 128 Mio = 16 777 216. 128 Mio = 134 217 728 octets."),
    ("Une ligne brute", "Une seule mesure. Ne pas utiliser une seule ligne, le minimum ou seulement la première seed pour établir le classement."),
    ("Recherche", "26 configurations × 4 volumes × 3 seeds × 3 répétitions = 936 mesures. Sert à sélectionner les finalistes."),
    ("Confirmation", "11 configurations par volume × 4 volumes × 7 nouvelles seeds × 9 mesures = 2 772 mesures. Sert à confirmer les résultats."),
    ("Temps typique", "Pour chaque seed, médiane des 9 mesures (la cinquième une fois les temps triés). Puis médiane de ces 7 médianes (la quatrième une fois triées). Ce n'est pas la moyenne des 63 mesures."),
    ("Gain de débit", "Dans ce classeur, chaque moteur est comparé au même Montmory précédent : pallas_original_32. Rapport = temps précédent / temps moteur, calculé pour chaque seed ; on prend ensuite la médiane des 7 rapports. Rapport > 1 et gain > 0 = le moteur de la ligne est plus rapide que cette référence."),
    ("Pourcentages et débits", "La médiane de rapports n'est pas nécessairement le rapport des médianes. C'est pourquoi +23,25 % ne se retrouve pas exactement en divisant 89,49 par 73,50."),
    ("Temps économisé", "+23,25 % de débit correspond à environ 18,86 % de temps économisé selon le rapport apparié, et non à 23,25 % de temps économisé."),
    ("Attention au JSON d'origine", "primary_speedup_vs_engine désigne la vitesse du candidat primaire PAR RAPPORT au moteur de la ligne. La colonne de ce classeur est volontairement orientée autrement : vitesse DU MOTEUR DE LA LIGNE par rapport à Montmory précédent. Les deux colonnes ne doivent pas être confondues."),
    ("Moteurs Montmory", "Les noms jax_original, jax_mulhi, jax_carry et pallas_* sont des implémentations du même flux Montmory. carry désigne un calcul exact de retenue ; 64 ou 128 désigne le nombre de lignes de la tuile Pallas, pas le nombre de cœurs physiques. Les 2048 lanes logiques restent fixes."),
    ("Deux Threefry", "threefry2x32 = API JAX ; threefry_pallas_* = noyau Pallas à 20 tours compatible avec le même flux de sortie de cette API. Ils peuvent avoir des vitesses très différentes."),
    ("RBG et témoin", "RBG est seulement contrôlé pour la répétabilité sur le même backend. write_NOT_PRNG est un programme d'écriture, pas un générateur aléatoire."),
    ("Seed", "Identifiant d'initialisation, pas une note de qualité. Conservé comme texte dans Excel pour éviter l'arrondi des grands entiers."),
    ("Bloc et répétition", "Trois blocs et trois répétitions dans chaque bloc en confirmation. Le classeur affiche 1, 2, 3 ; le CSV d'origine commence à 0."),
    ("Empreinte SHA256", "Identifie les octets de sortie, pas la vitesse ni la qualité statistique. Même seed, volume et flux Montmory : même empreinte attendue, quelle que soit la variante."),
    ("Configuration retenue", "Primaire choisi pendant la recherche, adopté seulement si son gain confirmé atteint 3 % et s'il est plus rapide sur les 7 seeds. Sinon conservation de pallas_original_32. Un secondaire plus rapide en confirmation n'est pas promu après avoir vu ces données."),
    ("Cas 64 Mio", "pallas_carry_64 affiche 63,99 Go/s, contre 62,46 pour le primaire pallas_carry_128. Le primaire reste retenu selon la règle prédéfinie ; ce choix n'affirme pas qu'il est le maximum de ce tableau."),
    ("Portée du temps", "Lancement, allocation et attente inclus. Compilation, préparation des clés, transfert des sorties et contrôle exclus. Ce n'est ni le temps du noyau seul ni le temps d'une application VaR complète."),
    ("Portée des validations", "Contrôles de conformité réussis et audit des empreintes ; aucune nouvelle batterie TestU01/TACM et aucune preuve d'optimalité globale. Les gains sont issus d'une session, pas de sessions indépendantes."),
    ("Données sources", "20261021T185102Z-tpu-opt/campaign/measurements.csv ; aucune modification du CSV d'origine."),
]
for row in guide_rows:
    guide.append(row)
guide.column_dimensions["A"].width = 29
guide.column_dimensions["B"].width = 110
guide.freeze_panes = "B2"
for row in guide:
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide.row_dimensions[row[0].row].height = max(34, 15*((len(str(row[1].value))//105)+1))
    row[0].font = Font(bold=True, color="154360")
guide.row_dimensions[1].height = 32
guide["A1"].font = Font(size=14, bold=True)


def table(sheet, headers, records, widths, name):
    sheet.append(headers)
    for row in records:
        sheet.append(row)
    sheet.freeze_panes = "C2"
    sheet.sheet_view.showGridLines = False
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for cell in sheet[1]:
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 45
    tab = Table(displayName=name, ref=f"A1:{get_column_letter(len(headers))}{len(records)+1}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(tab)


summary = wb.create_sheet("Synthèse")
headers = ["Volume", "Moteur", "Famille", "Implémentation", "Décision Montmory", "Mesures",
           "Temps typique (ms)", "Débit (Go/s)", "Rapport vs Montmory précédent",
           "Gain débit vs précédent (%)", "Seeds plus rapides / 7"]
records = []
for count in sizes:
    entries = sorted((r for r in audit["summary"] if r["words"] == count), key=lambda r: r["median_microseconds"])
    for row in entries:
        name = row["engine"]
        ratios = [median(count, "pallas_original_32", seed)/median(count, name, seed) for seed in report["confirm_seeds"]]
        records.append([labels[count], name, family(name), implementation(name), note(count, name), 63,
                        row["median_microseconds"]/1000, row["GB_per_second"], st.median(ratios),
                        100*(st.median(ratios)-1), sum(v>1 for v in ratios)])
table(summary, headers, records, [14,28,25,17,34,12,20,18,24,25,24], "SyntheseConfirmation")
for row in summary.iter_rows(min_row=2):
    for col in (7, 8, 9): row[col-1].number_format = "0.000"
    row[9].number_format = '0.00" %"'
    if row[4].value == "Montmory retenu":
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="D5F5E3")
            cell.font = Font(bold=True)

example = wb.create_sheet("Exemple 128 Mio", 1)
count = 16777216
names = ["pallas_original_32", "pallas_carry_64", "threefry2x32", "threefry_pallas_128"]
ex_records = []
for name in names:
    first = next(r for r in raw if r["phase"] == "confirmation" and r["words"] == count and
                 r["engine"] == name and r["seed"] == 7 and r["block"] == r["repetition"] == 0)
    seed_med = median(count, name, 7)
    typical = st.median(median(count, name, seed) for seed in report["confirm_seeds"])
    ex_records.append([name, first["nanoseconds"], first["nanoseconds"]/1e6, seed_med/1e6,
                       typical/1e6, 8*count/typical])
table(example, ["Moteur", "Première mesure seed 7 (ns)", "Cette mesure (ms)",
                "Médiane des 9 mesures seed 7 (ms)", "Temps typique des 7 seeds (ms)", "Débit typique (Go/s)"],
      ex_records, [29,28,23,35,34,26], "Exemple128Mio")
for row in example.iter_rows(min_row=2, max_row=5):
    for cell in row[2:]: cell.number_format = "0.000000"
example["A8"] = "Une mesure isolée peut inverser le classement."
example["A9"] = "Première mesure : Montmory carry64 = 1,364480 ms ; Threefry Pallas = 1,457230 ms. Montmory gagne CETTE mesure."
example["A10"] = "Sur les 9 mesures de la seed 7 : Montmory = 1,462309 ms ; Threefry Pallas = 1,457230 ms. Threefry est légèrement devant."
example["A11"] = "Temps typique sur les 7 seeds : Montmory = 1,499840 ms ; Threefry Pallas = 1,443110 ms. Threefry reste devant."
example["A13"] = "Les deux variantes Threefry fournissent le même flux de référence de cette API ; leur implémentation est différente."
for i in (8, 9, 10, 11, 13):
    example.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    example.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="center")
    example.row_dimensions[i].height = 30

seed_sheet = wb.create_sheet("Médianes par seed")
seed_records = []
for count in sizes:
    for seed in report["confirm_seeds"]:
        for name in report["selection"][str(count)]["finalists"]:
            med = median(count, name, seed)
            ref = median(count, "pallas_original_32", seed)
            seed_records.append([labels[count], name, str(seed), 9, med, med/1e6, 8*count/med, ref/med, 100*(ref/med-1)])
table(seed_sheet, ["Volume", "Moteur", "Seed (texte)", "Mesures", "Temps médian (ns)",
                   "Temps médian (ms)", "Débit (Go/s)", "Rapport vs précédent", "Gain débit vs précédent (%)"],
      seed_records, [14,28,25,12,22,23,18,23,27], "MedianesSeeds")
for row in seed_sheet.iter_rows(min_row=2):
    row[2].number_format = "@"
    for cell in row[5:8]: cell.number_format = "0.000000"
    row[8].number_format = '0.00" %"'

raw_sheet = wb.create_sheet("Mesures individuelles")
raw_records = [["Recherche" if r["phase"] == "search" else "Confirmation", labels[r["words"]],
                r["engine"], str(r["seed"]), r["block"]+1, r["repetition"]+1, r["words"],
                r["nanoseconds"], r["nanoseconds"]/1e6, 8*r["words"]/r["nanoseconds"],
                "Référence CPU fixe" if r["check"] == "fixed_cpu_reference" else "Même backend seulement",
                r["output_sha256"]] for r in raw]
table(raw_sheet, ["Phase", "Volume", "Moteur", "Seed (texte)", "Bloc (1 à 3)", "Répétition (1 à 3)",
                  "Mots de 64 bits", "Temps (ns)", "Temps (ms)", "Débit individuel (Go/s)", "Contrôle", "Empreinte SHA256"],
      raw_records, [18,14,29,25,16,21,21,20,19,27,28,72], "MesuresCompletes")
for row in raw_sheet.iter_rows(min_row=2):
    row[3].number_format = "@"
    row[8].number_format = "0.000000"
    row[9].number_format = "0.000"

wb.save(DEST)
# Re-open the actual deliverable and verify numeric values, all source rows and large seeds.
check = load_workbook(DEST, read_only=True, data_only=True)
assert check.sheetnames == ["Lire d'abord", "Exemple 128 Mio", "Synthèse", "Médianes par seed", "Mesures individuelles"]
assert check["Synthèse"].max_row == 45
assert check["Médianes par seed"].max_row == 309
assert check["Mesures individuelles"].max_row == 3709
for cells, expected in zip(check["Mesures individuelles"].iter_rows(min_row=2, values_only=True), raw_records, strict=True):
    assert isinstance(cells[3], str) and cells[3] == expected[3]
    for i in (0, 1, 2, 4, 5, 6, 7, 10, 11): assert cells[i] == expected[i]
    for i in (8, 9): assert abs(cells[i]-expected[i]) <= max(1, abs(expected[i]))*1e-12
for cells, expected in zip(check["Synthèse"].iter_rows(min_row=2, values_only=True), records, strict=True):
    for i in (0,1,2,3,4,5,10): assert cells[i] == expected[i]
    for i in (6,7,8,9): assert abs(cells[i]-expected[i]) <= max(1, abs(expected[i]))*1e-12
check.close()
csv_path = AUDIT / "synthese-lisible.csv"
with csv_path.open("w", newline="", encoding="utf-8-sig") as file:
    writer = csv.writer(file, delimiter=";")
    writer.writerow(headers)
    for row in records:
        writer.writerow([str(v).replace(".", ",") if isinstance(v, float) else v for v in row])
(AUDIT / "readable-workbook-checks.json").write_text(json.dumps(dict(status="pass", workbook=str(DEST),
    sheets=5, raw_rows_checked=len(raw_records), summary_rows_checked=len(records), seed_rows=308,
    seeds_preserved_as_text=True, source_csv_unchanged=True), indent=2)+"\n")
print(DEST)
print("PASS: 3708 source rows, 44 summary rows, large seeds preserved as text.")
